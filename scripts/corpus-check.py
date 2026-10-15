#!/usr/bin/env python3
"""
Open exported docx files with python-docx, a reader that is not ours, and report what it finds.

    python -I scripts/corpus-check.py FILE_OR_DIR [FILE_OR_DIR ...]

A file passes when it is a sound zip, every XML part in it is well formed, python-docx opens
it, and, where `<name>.expect.json` sits beside it, what python-docx sees agrees with the
expectations written there. The report is JSON on stdout, one object per file; the exit
status is 1 if any file failed.

The expectations file is optional and all of its keys are optional:

    {
      "texts":     ["a string that must occur in the document's text"],
      "headings":  [[1, "Heading 1 text"]],
      "tables":    [[rows, columns]],
      "images":    3,
      "comments":  2,
      "sections":  [{"orientation": "landscape", "paper": "Letter"}],
      "header":    "text that must occur in the first section's header",
      "footer":    "text that must occur in the first section's footer",
      "footnotes": 3,
      "tracked":   {"inserted": 2, "deleted": 2}
    }

Run it with `python -I` (isolated): the files are other people's. LibreOffice is not part of
this check: it is not installed here, so nothing in it says what the pages look like.
"""

import json
import sys
import zipfile
from pathlib import Path

from lxml import etree

try:
    from docx import Document
except ImportError:  # pragma: no cover
    print(json.dumps({"error": "python-docx is not installed in this interpreter"}))
    sys.exit(2)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

PAPERS = {  # portrait (short, long) in inches, as python-docx reports
    "A4": (8.27, 11.69),
    "Letter": (8.5, 11.0),
    "Legal": (8.5, 14.0),
    "A5": (5.83, 8.27),
}


def paper_of(width_in: float, height_in: float):
    short, long = sorted((width_in, height_in))
    for name, (s, l) in PAPERS.items():
        if abs(short - s) < 0.06 and abs(long - l) < 0.06:
            return name
    return None


def inspect(path: Path) -> dict:
    out: dict = {"file": path.name, "problems": []}
    problem = out["problems"].append

    # 1. The zip: sound, and every XML part parses.
    try:
        with zipfile.ZipFile(path) as z:
            bad = z.testzip()
            if bad:
                problem(f"corrupt zip entry {bad}")
            out["parts"] = len(z.namelist())
            for name in z.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    try:
                        etree.fromstring(z.read(name), etree.XMLParser(resolve_entities=False, no_network=True))
                    except etree.XMLSyntaxError as e:
                        problem(f"{name} is not well formed: {e}")
            if "[Content_Types].xml" not in z.namelist():
                problem("no [Content_Types].xml")
            names = set(z.namelist())
            rels = z.read("word/_rels/document.xml.rels").decode("utf8") if "word/_rels/document.xml.rels" in names else ""
            # Every internal relationship target must exist.
            for r in etree.fromstring(rels.encode("utf8")) if rels else []:
                if r.get("TargetMode") == "External":
                    continue
                target = r.get("Target", "")
                full = target.lstrip("/") if target.startswith("/") else "word/" + target
                parts = []
                for seg in full.split("/"):
                    if seg == "..":
                        parts and parts.pop()
                    elif seg not in ("", "."):
                        parts.append(seg)
                if "/".join(parts) not in names:
                    problem(f"relationship {r.get('Id')} points at {target}, which is not in the file")
    except zipfile.BadZipFile as e:
        problem(f"not a zip: {e}")
        return out

    # 2. python-docx opens it.
    try:
        doc = Document(str(path))
    except Exception as e:  # noqa: BLE001 - any failure is the finding
        problem(f"python-docx could not open it: {type(e).__name__}: {e}")
        return out

    paragraphs = doc.paragraphs
    out["paragraphs"] = len(paragraphs)
    out["headings"] = [
        [int(p.style.name.split()[-1]) if p.style.name.split()[-1].isdigit() else 0, p.text]
        for p in paragraphs
        if p.style is not None and p.style.name.lower().startswith("heading")
    ]
    out["tables"] = [[len(t.rows), len(t.columns)] for t in doc.tables]
    out["images"] = len(doc.inline_shapes)
    text = "\n".join(p.text for p in paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    out["text_length"] = len(text)

    sections = []
    for s in doc.sections:
        w, h = s.page_width.inches, s.page_height.inches
        sections.append(
            {
                "orientation": "landscape" if w > h else "portrait",
                "paper": paper_of(w, h),
                "margins_in": [round(x.inches, 2) for x in (s.top_margin, s.right_margin, s.bottom_margin, s.left_margin)],
                "different_first": bool(s.different_first_page_header_footer),
            }
        )
    out["sections"] = sections
    first = doc.sections[0]
    out["header"] = "\n".join(p.text for p in first.header.paragraphs)
    out["footer"] = "\n".join(p.text for p in first.footer.paragraphs)
    out["footer_fields"] = [i.text.strip() for i in first.footer._element.iter(W + "instrText")]

    comments = getattr(doc, "comments", None)
    out["comments"] = len(list(comments)) if comments is not None else None

    body = doc.element.body
    out["footnote_refs"] = len(list(body.iter(W + "footnoteReference")))
    out["tracked"] = {"inserted": len(list(body.iter(W + "ins"))), "deleted": len(list(body.iter(W + "del")))}

    # 3. What the file was supposed to hold.
    expect_path = path.with_suffix(".expect.json")
    if expect_path.exists():
        expect = json.loads(expect_path.read_text("utf8"))
        for s in expect.get("texts", []):
            if s not in text:
                problem(f"text {s!r} is not in the document")
        for lvl, s in expect.get("headings", []):
            if [lvl, s] not in out["headings"]:
                problem(f"heading {lvl} {s!r} is missing; found {out['headings']}")
        if "tables" in expect and expect["tables"] != out["tables"]:
            problem(f"tables are {out['tables']}, expected {expect['tables']}")
        if "images" in expect and expect["images"] != out["images"]:
            problem(f"{out['images']} inline pictures, expected {expect['images']}")
        if "comments" in expect and expect["comments"] != out["comments"]:
            problem(f"{out['comments']} comments, expected {expect['comments']}")
        for i, want in enumerate(expect.get("sections", [])):
            got = sections[i] if i < len(sections) else {}
            for k, v in want.items():
                if got.get(k) != v:
                    problem(f"section {i} {k} is {got.get(k)!r}, expected {v!r}")
        if "header" in expect and expect["header"] not in out["header"]:
            problem(f"header {out['header']!r} lacks {expect['header']!r}")
        if "footer" in expect and expect["footer"] not in out["footer"]:
            problem(f"footer {out['footer']!r} lacks {expect['footer']!r}")
        if "footnotes" in expect and expect["footnotes"] != out["footnote_refs"]:
            problem(f"{out['footnote_refs']} footnote references, expected {expect['footnotes']}")
        if "tracked" in expect and expect["tracked"] != out["tracked"]:
            problem(f"tracked changes are {out['tracked']}, expected {expect['tracked']}")
    return out


def main(argv) -> int:
    targets = []
    for a in argv:
        p = Path(a)
        targets += sorted(p.glob("*.docx")) if p.is_dir() else [p]
    if not targets:
        print(json.dumps({"error": "no files given"}))
        return 2
    reports = [inspect(t) for t in targets]
    print(json.dumps(reports, indent=2, ensure_ascii=False))
    return 1 if any(r["problems"] for r in reports) else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
