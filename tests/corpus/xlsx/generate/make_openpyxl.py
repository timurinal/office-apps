#!/usr/bin/env python3
"""
The generated half of the xlsx corpus: one workbook for each family of features, made with
openpyxl, which writes what Excel writes for the things below and nothing it has not been told.

    python -I make_openpyxl.py OUT_DIR

Each `NN-name.xlsx` comes with `NN-name.expect.json`, which says what is true of the workbook
independently of this repository's code: the value every cell should come to (worked out here, in
Python, by hand), the cells that hold formulas, and the formats, sizes and rules a reader should find.
openpyxl does not calculate, so the values a formula comes to are written into the file beside it
(`inject_cached`), as Excel does; the reader then has two things to agree with, the file's cache and
this script's arithmetic.

The files are committed so that the tests depend on neither this script nor openpyxl. Run it again
only to change them.
"""

import json
import re
import sys
import zipfile
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule, Rule
from openpyxl.styles import Alignment, Border, Color, Font, GradientFill, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ERRORS = {"#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!"}


def encode(v):
    """A value as the expectations file says it."""
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, (datetime, date, time)):
        return to_excel(v)
    if isinstance(v, str) and v in ERRORS:
        return {"error": v}
    return v


class Book:
    """A workbook being made, and what is true of it."""

    def __init__(self, name):
        self.name = name
        self.wb = openpyxl.Workbook()
        self.cached = {}  # sheet title -> {ref: value}
        self.values = {}
        self.formulas = []
        self.extra = {}

    def sheet(self, title=None, first=False):
        if first:
            ws = self.wb.active
            ws.title = title
        else:
            ws = self.wb.create_sheet(title)
        return ws

    def put(self, ws, ref, value, fmt=None, expect=True):
        ws[ref] = value
        if isinstance(value, str) and not value.startswith("="):
            ws[ref].data_type = "s"
        if fmt:
            ws[ref].number_format = fmt
        if expect:
            self.values[f"{ws.title}!{ref}"] = encode(value)

    def text(self, ws, ref, value, expect=True):
        """A string that is a string, whatever it starts with."""
        ws[ref].value = value
        ws[ref].data_type = "s"
        if expect:
            self.values[f"{ws.title}!{ref}"] = value

    def error(self, ws, ref, code):
        ws[ref].value = code
        ws[ref].data_type = "e"
        self.values[f"{ws.title}!{ref}"] = {"error": code}

    def formula(self, ws, ref, text, cached, fmt=None):
        ws[ref] = "=" + text
        if fmt:
            ws[ref].number_format = fmt
        self.cached.setdefault(ws.title, {})[ref] = cached
        self.values[f"{ws.title}!{ref}"] = encode(cached)
        self.formulas.append(f"{ws.title}!{ref}")

    def save(self, out):
        path = Path(out) / f"{self.name}.xlsx"
        self.wb.save(path)
        inject_cached(path, self.wb, self.cached)
        expect = {"values": self.values, "formulas": sorted(self.formulas), **self.extra}
        (Path(out) / f"{self.name}.expect.json").write_text(json.dumps(expect, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")


def inject_cached(path, wb, cached):
    """Put each formula's value beside it, as Excel does. openpyxl writes `<f>…</f><v />`."""
    if not cached:
        return
    sheet_files = {ws.title: f"xl/worksheets/sheet{i + 1}.xml" for i, ws in enumerate(wb.worksheets)}
    zin = zipfile.ZipFile(path)
    parts = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    pat = re.compile(r'<c r="([A-Z]+[0-9]+)"((?: s="[0-9]+")?)><f>(.*?)</f><v ?/></c>')
    for title, cells in cached.items():
        name = sheet_files[title]
        xml = parts[name].decode("utf-8")
        seen = set()

        def put(m):
            ref, style, formula = m.group(1), m.group(2), m.group(3)
            if ref not in cells:
                return m.group(0)
            seen.add(ref)
            value = cells[ref]
            if isinstance(value, bool):
                return '<c r="%s"%s t="b"><f>%s</f><v>%d</v></c>' % (ref, style, formula, 1 if value else 0)
            if isinstance(value, (datetime, date, time)):
                return '<c r="%s"%s><f>%s</f><v>%s</v></c>' % (ref, style, formula, repr(float(to_excel(value))))
            if isinstance(value, (int, float)):
                return '<c r="%s"%s><f>%s</f><v>%s</v></c>' % (ref, style, formula, repr(value))
            if isinstance(value, str) and value in ERRORS:
                return '<c r="%s"%s t="e"><f>%s</f><v>%s</v></c>' % (ref, style, formula, value)
            esc = str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            return '<c r="%s"%s t="str"><f>%s</f><v>%s</v></c>' % (ref, style, formula, esc)

        xml = pat.sub(put, xml)
        missing = set(cells) - seen
        assert not missing, f"{title}: formula cells not found: {sorted(missing)[:5]}"
        parts[name] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in parts.items():
            zout.writestr(n, data)


# ── 01 values ───────────────────────────────────────────────────────────────

def values(out):
    b = Book("01-values")
    ws = b.sheet("Values", first=True)
    for i, n in enumerate([0, 1, -1, 42, 3.14159, 1e-7, 1.5e12, 123456789012345, -0.5, 0.1 + 0.2, 1 / 3, 255, 1e21, 2 ** 53]):
        b.put(ws, f"A{i + 1}", n)
    strings = ["plain", "Ünïcödé — ✓ ☕", "日本語のテキスト", "  leading and trailing  ", "line1\nline2", "x" * 1000, "007", "'quoted", "a, b; c", 'say "hi"', "<tag> & more", "שלום", "emoji 😀🚀"]
    for i, s in enumerate(strings):
        b.text(ws, f"B{i + 1}", s)
    b.text(ws, f"B{len(strings) + 1}", "=not a formula")
    b.put(ws, "C1", True)
    b.put(ws, "C2", False)
    dates = [
        (date(2024, 3, 5), "yyyy-mm-dd"),
        (datetime(2024, 3, 5, 14, 30, 15), "yyyy-mm-dd hh:mm:ss"),
        (time(8, 0), "h:mm AM/PM"),
        (date(1900, 3, 1), "dd/mm/yyyy"),
        (date(2999, 12, 31), "d mmmm yyyy"),
        (datetime(2000, 1, 1, 0, 0, 1), "mm/dd/yy hh:mm"),
    ]
    for i, (d, fmt) in enumerate(dates):
        b.put(ws, f"D{i + 1}", d, fmt=fmt)
    b.error(ws, "E1", "#N/A")
    b.error(ws, "E2", "#DIV/0!")
    b.error(ws, "E3", "#REF!")
    b.text(ws, "F1", "123")
    b.text(ws, "F2", "1.5")
    b.text(ws, "F3", "TRUE")
    b.text(ws, "F4", "2024-03-05")
    b.save(out)


# ── 02 formulas ─────────────────────────────────────────────────────────────

def formulas(out):
    b = Book("02-formulas")
    c = b.sheet("Calc", first=True)
    o = b.sheet("Other Sheet")
    q = b.sheet("Q1 '24")
    nums = [10, 20, 30, 40, 50]
    for i, n in enumerate(nums):
        b.put(c, f"A{i + 1}", n)
    for i, s in enumerate(["apple", "Banana", "cherry", "apple", "Date"]):
        b.text(c, f"B{i + 1}", s)
    b.put(o, "A1", 5)
    b.put(o, "A2", 7)
    b.put(q, "A1", 100)
    b.put(q, "B1", 0.2)
    f = b.formula
    f(c, "D1", "SUM(A1:A5)", 150)
    f(c, "D2", "AVERAGE(A1:A5)", 30)
    f(c, "D3", "MIN(A1:A5)+MAX(A1:A5)", 60)
    f(c, "D4", "COUNT(A1:B5)", 5)
    f(c, "D5", "COUNTA(A1:B5)", 10)
    f(c, "D6", "A1+A2*A3-A4/A5", 10 + 20 * 30 - 40 / 50)
    f(c, "D7", "(A1+A2)*(A3-A4)", (10 + 20) * (30 - 40))
    f(c, "D8", "A2^2", 400)
    f(c, "D9", "-A1+10%", -10 + 0.1)
    f(c, "E1", 'IF(A1>5,"big","small")', "big")
    f(c, "E2", "IF(AND(A1>5,A2<10),1,0)", 0)
    f(c, "E3", "IF(OR(A1>50,A2=20),TRUE,FALSE)", True)
    f(c, "E4", 'IF(NOT(A1=10),"no","yes")', "yes")
    f(c, "E5", 'IF(A1>100,"a",IF(A1>5,"b","c"))', "b")
    f(c, "E6", "IFERROR(1/0,-1)", -1)
    f(c, "E7", '_xlfn.IFS(A1>100,"a",A1>5,"b",TRUE,"c")', "b")
    f(c, "E8", '_xlfn.SWITCH(A2,10,"ten",20,"twenty","other")', "twenty")
    f(c, "E9", '_xlfn.TEXTJOIN("-",TRUE,B1:B3)', "apple-Banana-cherry")
    f(c, "E10", '_xlfn.CONCAT(B1,"/",B2)', "apple/Banana")
    f(c, "F1", 'VLOOKUP(30,A1:B5,2,FALSE)', "cherry")
    f(c, "F2", "INDEX(A1:A5,MATCH(40,A1:A5,0))", 40)
    f(c, "F3", 'SUMIF(B1:B5,"apple",A1:A5)', 50)
    f(c, "F4", 'COUNTIF(B1:B5,"a*")', 2)
    f(c, "F5", 'SUMIFS(A1:A5,B1:B5,"apple",A1:A5,">10")', 40)
    f(c, "F6", "MATCH(25,A1:A5,1)", 2)
    f(c, "G1", 'CONCATENATE(B1," ",B2)', "apple Banana")
    f(c, "G2", 'LEFT(B3,3)&RIGHT(B3,2)&MID(B3,2,2)', "cheyhhe"[:0] + "che" + "ry" + "he")
    f(c, "G3", "LEN(B2)", 6)
    f(c, "G4", 'UPPER(B1)&LOWER(B2)', "APPLEbanana")
    f(c, "G5", 'SUBSTITUTE(B1,"p","P")', "aPPle")
    f(c, "G6", 'TRIM("  a   b  ")', "a b")
    f(c, "G7", 'TEXT(1234.5,"#,##0.00")', "1,234.50")
    f(c, "G8", 'TEXT(0.256,"0.0%")', "25.6%")
    f(c, "H1", "ROUND(2.567,1)", 2.6)
    f(c, "H2", "ABS(-3.5)+MOD(10,3)", 4.5)
    f(c, "H3", "POWER(2,10)+SQRT(81)", 1033)
    f(c, "H4", "INT(-2.5)+CEILING(2.1,1)", -3 + 3)
    f(c, "H5", "PRODUCT(A1:A3)", 6000)
    f(c, "H6", "ROUNDUP(2.1,0)+ROUNDDOWN(2.9,0)", 5)
    f(c, "I1", "DATE(2024,3,5)", date(2024, 3, 5), fmt="yyyy-mm-dd")
    f(c, "I2", "YEAR(I1)*10000+MONTH(I1)*100+DAY(I1)", 20240305)
    f(c, "I3", "I1+30", date(2024, 4, 4), fmt="yyyy-mm-dd")
    f(c, "I4", "EDATE(I1,2)", date(2024, 5, 5), fmt="yyyy-mm-dd")
    f(c, "I5", "WEEKDAY(I1)", 3)
    f(c, "I6", "DATE(2024,1,31)-DATE(2023,12,25)", 37)
    f(c, "J1", "'Other Sheet'!A1+'Other Sheet'!A2", 12)
    f(c, "J2", "SUM('Other Sheet'!A1:A2)*2", 24)
    f(c, "J3", "'Q1 ''24'!A1*'Q1 ''24'!B1", 20)
    f(c, "J4", "$A$1+A$2+$A3", 10 + 20 + 30)
    f(c, "J5", "SUM(A:A)", 150)
    f(c, "J6", "SUM('Other Sheet'!1:1)", 5 + 155)
    f(o, "B1", "Calc!D1+A1", 155)
    f(o, "B2", "SUM(Calc!A1:A5)/COUNT(Calc!A1:A5)", 30)
    # A running total down a column: the same formula, copied, as a person fills it.
    for i in range(1, 6):
        f(c, f"K{i}", f"A{i}" if i == 1 else f"K{i - 1}+A{i}", sum(nums[:i]))
    # Errors the formulas come to.
    f(c, "L1", "1/0", "#DIV/0!")
    f(c, "L2", "NA()", "#N/A")
    f(c, "L3", 'VALUE("abc")', "#VALUE!")
    f(c, "L4", "SQRT(-1)", "#NUM!")
    # Functions this engine does not have: the file's last value is kept. (openpyxl's theme is Office 2007's, so Accent 1 is 4F81BD.)
    f(c, "M1", "_xlfn.LAMBDA(x,x*2)(21)", 42)
    f(c, "M2", "SUM(INDIRECT(\"A1:A3\"))", 60)
    b.extra["kept_as_values"] = ["Calc!M1", "Calc!M2"]
    b.save(out)


# ── 03 formats ──────────────────────────────────────────────────────────────

def formats(out):
    b = Book("03-formats")
    ws = b.sheet("Formats", first=True)
    styles = {}
    fmts = [
        ("0.00", 1234.5),
        ("#,##0", 1234567),
        ("0%", 0.256),
        ("0.00E+00", 12345.678),
        ('"$"#,##0.00', 1234.5),
        ("[Red]-0.00;0.00", -5.5),
        ("yyyy-mm-dd", date(2024, 3, 5)),
        ("dd/mm/yyyy hh:mm", datetime(2024, 3, 5, 14, 30)),
        ("h:mm AM/PM", time(14, 30)),
        ("# ?/?", 0.75),
        ('"£"#,##0;[Red]-"£"#,##0', -1234),
        ('0.0" kg"', 72.5),
        ("@", "text"),
    ]
    for i, (fmt, v) in enumerate(fmts):
        ref = f"A{i + 1}"
        b.put(ws, ref, v, fmt=fmt)
        styles[f"Formats!{ref}"] = {"numFmt": fmt}
    row = len(fmts) + 2
    # Fonts.
    fonts = [
        ("B", Font(name="Arial", size=14, bold=True), {"font": "Arial", "size": 14, "bold": True}),
        ("C", Font(italic=True, color="FFC0392B"), {"italic": True, "color": "#c0392b"}),
        ("D", Font(underline="single", strike=True), {"underline": True, "strike": True}),
        ("E", Font(name="Courier New", size=9, color=Color(theme=4, tint=0.3999755851924192)), {"font": "Courier New", "size": 9, "color": "near:#95b3d7"}),
        ("F", Font(name="Calibri", size=11, color="FF000000"), {}),
    ]
    for i, (col, font, expect) in enumerate(fonts):
        ref = f"{col}{row}"
        b.put(ws, ref, f"font {i}")
        ws[ref].font = font
        styles[f"Formats!{ref}"] = expect
    row += 2
    fills = [
        ("B", PatternFill("solid", fgColor="FFFFF2CC"), {"fill": "#fff2cc"}),
        ("C", PatternFill("solid", fgColor=Color(theme=5, tint=0.5999938962981048)), {"fill": "near:#e6b9b8"}),
        ("D", PatternFill("solid", fgColor=Color(indexed=13)), {"fill": "#ffff00"}),
        ("E", PatternFill("lightGray", fgColor="FFFF0000"), {}),
        ("F", GradientFill(stop=("FF0000", "0000FF")), {}),
    ]
    for i, (col, fill, expect) in enumerate(fills):
        ref = f"{col}{row}"
        b.put(ws, ref, f"fill {i}")
        ws[ref].fill = fill
        styles[f"Formats!{ref}"] = expect
    row += 2
    sides = ["thin", "medium", "thick", "dashed", "dotted", "double", "hair", "mediumDashed", "dashDot"]
    want = {"thin": "thin", "medium": "medium", "thick": "thick", "dashed": "dashed", "dotted": "dotted", "double": "double", "hair": "dotted", "mediumDashed": "medium", "dashDot": "dashed"}
    for i, s in enumerate(sides):
        ref = f"{get_column_letter(2 + i)}{row}"
        b.put(ws, ref, s)
        side = Side(style=s, color="FF123456")
        ws[ref].border = Border(top=side, bottom=side, left=Side(style=s), right=Side(style=s))
        styles[f"Formats!{ref}"] = {
            "borderTop": {"line": want[s], "color": "#123456"},
            "borderBottom": {"line": want[s], "color": "#123456"},
            "borderLeft": {"line": want[s]},
            "borderRight": {"line": want[s]},
        }
    row += 2
    aligns = [
        ("B", Alignment(horizontal="left", vertical="top"), {"h": "left", "v": "top"}),
        ("C", Alignment(horizontal="center", vertical="center", wrap_text=True), {"h": "center", "wrap": True}),  # centred vertically is the grid's own default, so it is not described
        ("D", Alignment(horizontal="right", vertical="bottom", indent=3), {"h": "right", "v": "bottom", "indent": 3}),
        ("E", Alignment(horizontal="justify"), {"h": "left"}),
        ("F", Alignment(text_rotation=45, shrink_to_fit=True), {}),
    ]
    for i, (col, al, expect) in enumerate(aligns):
        ref = f"{col}{row}"
        b.put(ws, ref, f"align {i}")
        ws[ref].alignment = al
        styles[f"Formats!{ref}"] = expect
    row += 2
    # A formatted column and a formatted row; a cell in each that says its own thing.
    ws.column_dimensions["H"].fill = PatternFill("solid", fgColor="FFE8F5E9")
    ws.column_dimensions["H"].font = Font(bold=True)
    b.put(ws, f"H{row}", "in a formatted column")
    ws[f"H{row}"].fill = PatternFill("solid", fgColor="FFE8F5E9")
    ws[f"H{row}"].font = Font(bold=True)
    b.put(ws, f"H{row + 1}", "with a format of its own")
    ws[f"H{row + 1}"].number_format = "0.0"
    ws.row_dimensions[row + 3].fill = PatternFill("solid", fgColor="FFE3F2FD")
    b.put(ws, f"B{row + 3}", "in a formatted row")
    ws[f"B{row + 3}"].fill = PatternFill("solid", fgColor="FFE3F2FD")
    styles[f"Formats!H{row}"] = {}
    styles[f"Formats!H{row + 1}"] = {"numFmt": "0.0"}
    styles[f"Formats!B{row + 3}"] = {}
    b.extra["styles"] = styles
    b.extra["column_styles"] = {"Formats!H": {"fill": "#e8f5e9", "bold": True}}
    b.extra["row_styles"] = {f"Formats!{row + 3}": {"fill": "#e3f2fd"}}
    b.extra["drops"] = ["gradient"]
    b.save(out)


# ── 04 layout ───────────────────────────────────────────────────────────────

def layout(out):
    b = Book("04-layout")
    ws = b.sheet("Sizes", first=True)
    for r in range(1, 12):
        for c in range(1, 8):
            b.put(ws, f"{get_column_letter(c)}{r}", r * 10 + c)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["C"].width = 9
    ws.column_dimensions["E"].hidden = True
    ws.column_dimensions["F"].width = 40
    ws.row_dimensions[2].height = 60
    ws.row_dimensions[3].hidden = True
    ws.row_dimensions[5].height = 15  # Excel's own default: says nothing
    ws.row_dimensions[6].height = 36
    ws.merge_cells("A8:C8")
    ws.merge_cells("D9:D11")
    ws.merge_cells("F8:G9")
    # openpyxl drops what a merge covers, so the file holds nothing there, and nothing is expected.
    for gone in ("B8", "C8", "D10", "D11", "F9", "G8", "G9"):
        del b.values[f"Sizes!{gone}"]
    thin = Side(style="medium")
    for ref, b_ in [("A8", Border(top=thin, left=thin, bottom=thin)), ("B8", Border(top=thin, bottom=thin)), ("C8", Border(top=thin, right=thin, bottom=thin))]:
        ws[ref].border = b_
    ws.sheet_properties.tabColor = "FF22C55E"
    r = b.sheet("Rows frozen")
    c = b.sheet("Columns frozen")
    both = b.sheet("Both frozen")
    for w in (r, c, both):
        for i in range(1, 6):
            b.put(w, f"A{i}", i)
            b.put(w, f"B{i}", i * 2)
    r.freeze_panes = "A3"
    c.freeze_panes = "C1"
    both.freeze_panes = "B2"
    r.sheet_properties.tabColor = Color(theme=4)
    h = b.sheet("Hidden sheet")
    b.put(h, "A1", "secret")
    h.sheet_state = "hidden"
    last = b.sheet("Last")
    b.put(last, "A1", "last")
    b.extra["layout"] = {
        "Sizes": {
            "widths": {"0": 140, "2": 63, "5": 280},
            "hiddenCols": [4],
            "heights": {"1": 80, "5": 48},
            "hiddenRows": [2],
            "merges": ["A8:C8", "D9:D11", "F8:G9"],
            "tab": "#22c55e",
            "frozen": [0, 0],
        },
        "Rows frozen": {"frozen": [2, 0], "tab": "near:#4f81bd"},
        "Columns frozen": {"frozen": [0, 2]},
        "Both frozen": {"frozen": [1, 1]},
        "Hidden sheet": {"hidden": True},
    }
    b.extra["order"] = ["Sizes", "Rows frozen", "Columns frozen", "Both frozen", "Hidden sheet", "Last"]
    b.save(out)


# ── 05 names ────────────────────────────────────────────────────────────────

def names(out):
    b = Book("05-names")
    data = b.sheet("Data", first=True)
    other = b.sheet("Other")
    for i in range(1, 6):
        b.put(data, f"A{i}", i * 10)
        b.put(other, f"A{i}", i)
    b.put(data, "C1", 0.25)
    wb = b.wb
    wb.defined_names["Rate"] = DefinedName("Rate", attr_text="Data!$C$1")
    wb.defined_names["Prices"] = DefinedName("Prices", attr_text="Data!$A$1:$A$5")
    wb.defined_names["Tax"] = DefinedName("Tax", attr_text="0.2")
    wb.defined_names["Double"] = DefinedName("Double", attr_text="Data!$A$1*2")
    wb.defined_names["Elsewhere"] = DefinedName("Elsewhere", attr_text="Other!$A$1:$A$5")
    data.defined_names["Local"] = DefinedName("Local", attr_text="Data!$A$2")
    other.defined_names["Local"] = DefinedName("Local", attr_text="Other!$A$2")
    f = b.formula
    f(data, "E1", "SUM(Prices)", 150)
    f(data, "E2", "Rate*100", 25)
    f(data, "E3", "A1*Tax", 2)
    f(data, "E4", "Double+1", 21)
    f(data, "E5", "SUM(Elsewhere)", 15)
    f(data, "E6", "Local", 20)
    f(other, "C1", "Local", 2)
    b.extra["names"] = [
        {"name": "Rate", "scope": None, "text": "Data!$C$1"},
        {"name": "Prices", "scope": None, "text": "Data!$A$1:$A$5"},
        {"name": "Tax", "scope": None, "text": "0.2"},
        {"name": "Double", "scope": None, "text": "Data!$A$1*2"},
        {"name": "Elsewhere", "scope": None, "text": "Other!$A$1:$A$5"},
        {"name": "Local", "scope": "Data", "text": "Data!$A$2"},
        {"name": "Local", "scope": "Other", "text": "Other!$A$2"},
    ]
    b.save(out)


# ── 06 validation ───────────────────────────────────────────────────────────

def validation(out):
    b = Book("06-validation")
    ws = b.sheet("Rules", first=True)
    lists = b.sheet("Lists")
    for i, v in enumerate(["north", "south", "east", "west"]):
        b.text(lists, f"A{i + 1}", v)
    for i, v in enumerate(["x", "y"]):
        b.text(ws, f"H{i + 1}", v)

    def add(dv, rng):
        ws.add_data_validation(dv)
        for part in rng.split():
            dv.add(part)

    add(DataValidation(type="list", formula1='"red,green,blue"', allow_blank=True, showErrorMessage=True, error="Pick a colour.", errorTitle="Colour"), "A2:A10")
    add(DataValidation(type="list", formula1="$H$1:$H$2", allow_blank=False), "B2:B10")
    add(DataValidation(type="list", formula1="Lists!$A$1:$A$4", showErrorMessage=True), "C2:C10")
    add(DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True, showErrorMessage=True, errorStyle="warning"), "D2:D10")
    add(DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0.5", showErrorMessage=True), "E2:E10")
    add(DataValidation(type="date", operator="greaterThan", formula1="45000", showErrorMessage=True), "F2:F10")
    add(DataValidation(type="textLength", operator="lessThanOrEqual", formula1="5", showErrorMessage=True, errorStyle="information"), "G2:G10")
    add(DataValidation(type="custom", formula1="AND(I2>0,I2<100)", showErrorMessage=True), "I2:I10")
    add(DataValidation(type="list", formula1='"yes,no"', showErrorMessage=True), "J2:J1048576")
    add(DataValidation(type="whole", operator="notBetween", formula1="5", formula2="6", showErrorMessage=True), "K2:K4 K8")
    b.extra["validations"] = [
        {"ranges": ["A2:A10"], "check": {"kind": "list", "items": ["red", "green", "blue"]}, "strict": True, "message": "Pick a colour."},
        {"ranges": ["B2:B10"], "check": {"kind": "listRange", "sheet": "Rules", "range": "H1:H2"}, "allowBlank": False},
        {"ranges": ["C2:C10"], "check": {"kind": "listRange", "sheet": "Lists", "range": "A1:A4"}, "strict": True, "allowBlank": False},
        {"ranges": ["D2:D10"], "check": {"kind": "integer", "op": "between", "a": 1, "b": 10}},
        {"ranges": ["E2:E10"], "check": {"kind": "number", "op": "ge", "a": 0.5}, "strict": True, "allowBlank": False},
        {"ranges": ["F2:F10"], "check": {"kind": "date", "op": "gt", "a": 45000}, "strict": True, "allowBlank": False},
        {"ranges": ["G2:G10"], "check": {"kind": "textLength", "op": "le", "a": 5}, "allowBlank": False},
        {"ranges": ["I2:I10"], "check": {"kind": "formula", "f": "=AND(I2>0,I2<100)"}, "strict": True, "allowBlank": False},
        {"ranges": ["J2:J1000"], "check": {"kind": "list", "items": ["yes", "no"]}, "strict": True, "allowBlank": False},
        {"ranges": ["K2:K4"], "check": {"kind": "integer", "op": "notBetween", "a": 5, "b": 6}, "strict": True, "allowBlank": False},
        {"ranges": ["K8"], "check": {"kind": "integer", "op": "notBetween", "a": 5, "b": 6}, "strict": True, "allowBlank": False},
    ]
    b.save(out)


# ── 07 conditional formats ──────────────────────────────────────────────────

def conditional(out):
    b = Book("07-conditional")
    ws = b.sheet("Rules", first=True)
    for r in range(1, 11):
        b.put(ws, f"A{r}", r)
        b.text(ws, f"B{r}", ["apple", "Banana", "cherry", "date", "apple", "fig", "grape", "Apple", "kiwi", "lime"][r - 1])
    b.put(ws, "D1", 3)
    red = DifferentialStyle(font=Font(bold=True, color="FF9C0006"), fill=PatternFill(bgColor="FFFFC7CE", fill_type="solid"))
    green = DifferentialStyle(font=Font(italic=True, strike=True, color="FF006100"), fill=PatternFill(bgColor="FFC6EFCE", fill_type="solid"))
    ws.conditional_formatting.add("A1:A10", CellIsRule(operator="greaterThan", formula=["7"], font=Font(bold=True, color="FF9C0006"), fill=PatternFill(bgColor="FFFFC7CE", fill_type="solid")))
    ws.conditional_formatting.add("A1:A10", CellIsRule(operator="between", formula=["2", "4"], font=Font(italic=True, color="FF006100")))
    ws.conditional_formatting.add("A1:A10", CellIsRule(operator="equal", formula=["$D$1"], fill=PatternFill(bgColor="FFFFEB9C", fill_type="solid")))
    ws.conditional_formatting.add("B1:B10", Rule(type="containsText", operator="containsText", text="an", dxf=red, formula=['NOT(ISERROR(SEARCH("an",B1)))']))
    ws.conditional_formatting.add("B1:B10", Rule(type="beginsWith", operator="beginsWith", text="a", dxf=green, formula=['LEFT(B1,LEN("a"))="a"']))
    ws.conditional_formatting.add("B1:B10", Rule(type="endsWith", operator="endsWith", text="e", dxf=green, formula=['RIGHT(B1,LEN("e"))="e"']))
    ws.conditional_formatting.add("B1:B10", Rule(type="duplicateValues", dxf=red))
    ws.conditional_formatting.add("C1:C10", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=PatternFill(bgColor="FFEEEEEE", fill_type="solid")))
    ws.conditional_formatting.add("E1:E10", Rule(type="top10", rank=3, dxf=red))
    ws.conditional_formatting.add("E1:E10", Rule(type="top10", rank=20, percent=True, bottom=True, dxf=green))
    ws.conditional_formatting.add("F1:F10", Rule(type="aboveAverage", dxf=red))
    ws.conditional_formatting.add("F1:F10", Rule(type="aboveAverage", aboveAverage=False, dxf=green))
    ws.conditional_formatting.add("G1:G10", ColorScaleRule(start_type="min", start_color="FFFFFFFF", end_type="max", end_color="FF63BE7B"))
    ws.conditional_formatting.add("H1:H10", ColorScaleRule(start_type="min", start_color="FFF8696B", mid_type="percentile", mid_value=50, mid_color="FFFFEB84", end_type="max", end_color="FF63BE7B"))
    ws.conditional_formatting.add("I1:I10", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
    ws.conditional_formatting.add("J1:J10", IconSetRule("3Arrows", "percent", [0, 33, 67]))
    ws.conditional_formatting.add("K1:K10", Rule(type="containsBlanks", dxf=red, formula=["LEN(TRIM(K1))=0"]))
    ws.conditional_formatting.add("L1:L10", Rule(type="containsErrors", dxf=red, formula=["ISERROR(L1)"]))
    ws.conditional_formatting.add("M1:M10", Rule(type="uniqueValues", dxf=green))
    for r in range(1, 11):
        for col, v in zip("EFGHIJK", [r * 7 % 11, r, r * r, r * 3, r, r, r if r % 3 else None]):
            if v is not None:
                b.put(ws, f"{col}{r}", v)
    b.extra["cf"] = {
        "Rules": [
            {"ranges": ["A1:A10"], "when": {"kind": "cell", "op": "gt", "a": 7}, "style": {"color": "#9c0006", "bold": True, "fill": "#ffc7ce"}},
            {"ranges": ["A1:A10"], "when": {"kind": "cell", "op": "between", "a": 2, "b": 4}, "style": {"color": "#006100", "italic": True}},
            {"ranges": ["A1:A10"], "when": {"kind": "formula", "f": "=A1=$D$1"}, "style": {"fill": "#ffeb9c"}},
            {"ranges": ["B1:B10"], "when": {"kind": "text", "op": "contains", "text": "an"}, "style": {"color": "#9c0006", "bold": True, "fill": "#ffc7ce"}},
            {"ranges": ["B1:B10"], "when": {"kind": "text", "op": "begins", "text": "a"}},
            {"ranges": ["B1:B10"], "when": {"kind": "text", "op": "ends", "text": "e"}},
            {"ranges": ["B1:B10"], "when": {"kind": "duplicate"}},
            {"ranges": ["C1:C10"], "when": {"kind": "formula", "f": "=MOD(ROW(),2)=0"}, "style": {"fill": "#eeeeee"}},
            {"ranges": ["E1:E10"], "when": {"kind": "top", "n": 3}},
            {"ranges": ["E1:E10"], "when": {"kind": "top", "n": 20, "percent": True, "bottom": True}},
            {"ranges": ["F1:F10"], "when": {"kind": "average"}},
            {"ranges": ["F1:F10"], "when": {"kind": "average", "below": True}},
            {"ranges": ["G1:G10"], "when": {"kind": "scale", "min": "#ffffff", "max": "#63be7b"}},
            {"ranges": ["H1:H10"], "when": {"kind": "scale", "min": "#f8696b", "mid": "#ffeb84", "max": "#63be7b"}},
            {"ranges": ["I1:I10"], "when": {"kind": "bar", "color": "#638ec6"}},
            {"ranges": ["K1:K10"], "when": {"kind": "blank"}},
            {"ranges": ["L1:L10"], "when": {"kind": "error"}},
            {"ranges": ["M1:M10"], "when": {"kind": "unique"}},
        ]
    }
    b.extra["left_out"] = ["icon set"]
    b.save(out)


# ── 08 large ────────────────────────────────────────────────────────────────

def large(out):
    b = Book("08-large")
    ws = b.sheet("Ledger", first=True)
    n = 20000
    running = 0.0
    total_c = 0.0
    for r in range(1, n + 1):
        c = (r % 97) * 1.25
        d = c * 2
        running += d
        total_c += c
        ws[f"A{r}"] = r
        ws[f"B{r}"] = f"Item {r}"
        ws[f"B{r}"].data_type = "s"
        ws[f"C{r}"] = c
        ws[f"D{r}"] = f"=C{r}*2"
        ws[f"E{r}"] = f"=D{r}" if r == 1 else f"=E{r - 1}+D{r}"
        ws[f"F{r}"] = date(2020, 1, 1 + (r % 28))
        ws[f"F{r}"].number_format = "dd/mm/yyyy"
        b.cached.setdefault("Ledger", {})[f"D{r}"] = d
        b.cached["Ledger"][f"E{r}"] = running
        if r in (1, 2, 97, 98, 10000, 19999, n):
            b.values[f"Ledger!C{r}"] = c
            b.values[f"Ledger!D{r}"] = d
            b.values[f"Ledger!E{r}"] = running
    ws[f"H1"] = f"=SUM(C1:C{n})"
    b.cached["Ledger"]["H1"] = total_c
    b.values["Ledger!H1"] = total_c
    b.extra["rows"] = n
    b.extra["large_formulas"] = n * 2 + 1
    # The formulas are many; the expectations name a few of them.
    b.formulas = ["Ledger!D1", "Ledger!E2", "Ledger!E20000", "Ledger!H1"]
    b.save(out)


# ── 09 names and text that are not ASCII ────────────────────────────────────

def unicode_names(out):
    b = Book("09-unicode-names")
    a = b.sheet("Données été", first=True)
    k = b.sheet("Año 2024")
    j = b.sheet("日本語")
    q = b.sheet("A&B (copy)")
    x = b.sheet("x y")
    b.put(a, "A1", 1)
    b.put(k, "A1", 2)
    b.put(j, "A1", 3)
    b.put(q, "A1", 4)
    b.put(x, "A1", 5)
    b.formula(a, "B1", "'Año 2024'!A1+'日本語'!A1+'A&B (copy)'!A1+'x y'!A1", 14)
    b.formula(k, "B1", "'Données été'!A1*10", 10)
    b.formula(x, "B1", "SUM('Données été'!A1,'Año 2024'!A1)", 3)
    b.text(a, "C1", "emoji 😀 and RTL שלום עולם and combining é")
    b.text(a, "C2", "y" * 32767)
    b.extra["order"] = ["Données été", "Año 2024", "日本語", "A&B (copy)", "x y"]
    b.save(out)


# ── 10 filter and table ─────────────────────────────────────────────────────

def filters(out):
    b = Book("10-filter-table")
    ws = b.sheet("Filtered", first=True)
    rows = [("name", "team", "score"), ("Ada", "red", 91), ("Ben", "blue", 78), ("Cy", "red", 64), ("Di", "green", 85), ("Ed", "blue", 52)]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            b.put(ws, f"{get_column_letter(c)}{r}", v)
    ws.auto_filter.ref = "A1:C6"
    ws.auto_filter.add_filter_column(1, ["red"])
    ws.row_dimensions[3].hidden = True
    ws.row_dimensions[5].hidden = True
    ws.row_dimensions[6].hidden = True
    t = b.sheet("Table")
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            b.put(t, f"{get_column_letter(c)}{r}", v)
    tab = Table(displayName="Scores", ref="A1:C6")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    t.add_table(tab)
    b.extra["filter"] = {"Filtered": "A1:C6", "hiddenRows": [2, 4, 5]}
    b.extra["table"] = "Table"
    b.save(out)


# ── 11 the 1904 date system ─────────────────────────────────────────────────

def date1904(out):
    b = Book("11-date1904")
    ws = b.sheet("Mac", first=True)
    b.wb.epoch = CALENDAR_MAC_1904
    b.put(ws, "A1", date(2024, 1, 1), fmt="yyyy-mm-dd")
    b.put(ws, "A2", datetime(2024, 3, 5, 12, 0), fmt="yyyy-mm-dd hh:mm")
    b.put(ws, "A3", date(2000, 2, 29), fmt="dd/mm/yyyy")
    b.put(ws, "B1", 5)
    # The file stores 1904-based serials; what is true is the date.
    b.values["Mac!A1"] = to_excel(date(2024, 1, 1))
    b.values["Mac!A2"] = to_excel(datetime(2024, 3, 5, 12, 0))
    b.values["Mac!A3"] = to_excel(date(2000, 2, 29))
    b.extra["date1904"] = True
    b.save(out)


# ── 12 what a file carries that is not read ─────────────────────────────────

def carried(out):
    b = Book("12-chart-links-notes")
    ws = b.sheet("Carried", first=True)
    for i in range(1, 6):
        b.put(ws, f"A{i}", i * 3)
        b.put(ws, f"B{i}", f"label {i}")
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=5))
    ws.add_chart(chart, "D2")
    b.text(ws, "C1", "web")
    ws["C1"].hyperlink = "https://example.com/"
    b.text(ws, "C2", "script")
    ws["C2"].hyperlink = "javascript:alert(1)"
    b.text(ws, "C3", "mail")
    ws["C3"].hyperlink = "mailto:ada@example.com"
    b.text(ws, "C4", "inside")
    ws["C4"].hyperlink = "#Carried!A1"
    ws["A2"].comment = Comment("a note on a cell", "Ada")
    b.extra["left_out"] = {"charts": 1, "links": 4, "notes": True}
    b.save(out)


def main():
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    for fn in (values, formulas, formats, layout, names, validation, conditional, large, unicode_names, filters, date1904, carried):
        fn(out)


if __name__ == "__main__":
    main()
