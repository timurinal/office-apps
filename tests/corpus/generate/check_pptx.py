#!/usr/bin/env python3
"""
Open exported presentations with python-pptx, a reader that is not ours, and report what it finds.

    python -I check_pptx.py DIR

For every `<name>.pptx` in DIR with a `<name>.check.json` beside it, the file must load, every
relationship in the package must point at a part that is there, every part must have a content type,
and what python-pptx sees must agree with the check file, which is what our own reader found in
the file that was exported: the slide size, the layouts, and for each slide its layout, whether it is
hidden, its notes, and each element's kind, placeholder, text, place, rotation and, by kind, the cells
and merges of a table, the type and values of a chart, the type and crop of a picture, the children of a
group. The report is JSON on stdout, one object per file; the exit status is 1 if any file has a problem.

Run it with `python -I` (isolated): the files, even our own, are not to be trusted to import code.
LibreOffice and PowerPoint are not part of this check: neither is installed on the machine this was written
on, so nothing here says what a slide looks like.
"""

import json
import posixpath
import sys
import warnings
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"

KIND = {
    "shape": {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER, MSO_SHAPE_TYPE.FREEFORM},
    "image": {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.PLACEHOLDER},
    "line": {MSO_SHAPE_TYPE.LINE},
    "table": {MSO_SHAPE_TYPE.TABLE},
    "chart": {MSO_SHAPE_TYPE.CHART},
    "group": {MSO_SHAPE_TYPE.GROUP},
}

PLACEHOLDER = {
    "title": {"TITLE"},
    "ctrTitle": {"CENTER_TITLE"},
    "subTitle": {"SUBTITLE"},
    "body": {"BODY", "OBJECT"},
    "pic": {"PICTURE"},
    "dt": {"DATE"},
    "ftr": {"FOOTER"},
    "sldNum": {"SLIDE_NUMBER"},
}

CHART = {
    "column": {"COLUMN_CLUSTERED"},
    "bar": {"BAR_CLUSTERED"},
    "line": {"LINE", "LINE_MARKERS"},
    "pie": {"PIE"},
    "scatter": {"XY_SCATTER", "XY_SCATTER_LINES", "XY_SCATTER_LINES_NO_MARKERS", "XY_SCATTER_SMOOTH", "XY_SCATTER_SMOOTH_NO_MARKERS"},
}


def package_problems(path: Path):
    """Every relationship has a target, every part a content type."""
    problems = []
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        ct = ET.fromstring(z.read("[Content_Types].xml"))
        defaults = {e.get("Extension").lower() for e in ct.findall(CT + "Default")}
        overrides = {e.get("PartName") for e in ct.findall(CT + "Override")}
        for n in names:
            if n == "[Content_Types].xml" or n.endswith("/"):
                continue
            ext = n.rsplit(".", 1)[-1].lower() if "." in n else ""
            if "/" + n not in overrides and ext not in defaults:
                problems.append(f"{n}: no content type")
        for o in overrides:
            if o.lstrip("/") not in names:
                problems.append(f"{o}: a content type for a part that is not there")
        for n in names:
            if not n.endswith(".rels"):
                continue
            base = posixpath.dirname(posixpath.dirname(n))
            for rel in ET.fromstring(z.read(n)).findall(PKG + "Relationship"):
                if rel.get("TargetMode") == "External":
                    continue
                target = rel.get("Target")
                full = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                if full not in names:
                    problems.append(f"{n}: {rel.get('Id')} points at {full}, which is not there")
    return problems


def close(a, b, tol):
    return abs(a - b) <= tol


def text_of(shape):
    if not shape.has_text_frame:
        return None
    t = shape.text_frame.text.replace("\x0b", "\n")
    return t if t.strip() else None


def compare(shape, want, tag, emu_per_px, problems):
    if shape.shape_type not in KIND[want["kind"]] and not (want["kind"] == "shape" and shape.is_placeholder):
        problems.append(f"{tag}: shape type {shape.shape_type} for a {want['kind']}")
        return
    if want.get("ph"):
        if not shape.is_placeholder:
            problems.append(f"{tag}: not a placeholder")
        elif shape.placeholder_format.type.name not in PLACEHOLDER[want["ph"]]:
            problems.append(f"{tag}: placeholder {shape.placeholder_format.type.name} for {want['ph']}")
    elif shape.is_placeholder:
        problems.append(f"{tag}: a placeholder that is not one")
    x, y, w, h = want["box"]
    got = [shape.left, shape.top, shape.width, shape.height]
    for g, e, label in zip(got, [x, y, w, h], "xywh"):
        if g is None or not close(g, e * emu_per_px, emu_per_px * 1.01):
            problems.append(f"{tag}: {label} {g} for {e * emu_per_px:.0f}")
    if want["kind"] != "group" and (want.get("rot") or 0) and abs((shape.rotation or 0) - want["rot"]) > 0.05:
        problems.append(f"{tag}: rotation {shape.rotation} for {want['rot']}")
    if want["kind"] in ("shape",):
        if text_of(shape) != want["text"]:
            problems.append(f"{tag}: text {text_of(shape)!r} for {want['text']!r}")
    if want["kind"] == "table":
        t = shape.table
        rows = [[c.text for c in r.cells] for r in t.rows]
        if len(rows) != want["table"]["rows"] or len(t.columns) != want["table"]["cols"]:
            problems.append(f"{tag}: table {len(rows)}x{len(t.columns)}")
        for r, row in enumerate(want["table"]["texts"]):
            for c, text in enumerate(row):
                if r < len(rows) and c < len(rows[r]) and rows[r][c] != text:
                    problems.append(f"{tag}: cell {r},{c} {rows[r][c]!r} for {text!r}")
        for s in want["table"]["spans"]:
            cell = t.cell(s["r"], s["c"])
            if s["merged"]:
                if not cell.is_spanned:
                    problems.append(f"{tag}: cell {s['r']},{s['c']} is not covered")
            elif not cell.is_merge_origin or cell.span_width != s["cs"] or cell.span_height != s["rs"]:
                problems.append(f"{tag}: cell {s['r']},{s['c']} spans {cell.span_width}x{cell.span_height} for {s['cs']}x{s['rs']}")
    if want["kind"] == "chart":
        chart = shape.chart
        if chart.chart_type.name not in CHART[want["chart"]["chart"]]:
            problems.append(f"{tag}: chart type {chart.chart_type.name} for {want['chart']['chart']}")
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        if want["chart"]["chart"] != "scatter" and cats != want["chart"]["categories"]:
            problems.append(f"{tag}: categories {cats} for {want['chart']['categories']}")
        series = [(s.name, [float(v) for v in s.values]) for s in plot.series]
        if series != [(s["name"], [float(v) for v in s["values"]]) for s in want["chart"]["series"]]:
            problems.append(f"{tag}: series {series}")
        try:
            blob = chart.part.chart_workbook.xlsx_part.blob
            if not blob.startswith(b"PK"):
                problems.append(f"{tag}: the chart's workbook is not a zip")
            else:
                workbook_problems(blob, want["chart"], tag, problems)
        except Exception as e:  # noqa: BLE001
            problems.append(f"{tag}: the chart has no workbook ({e})")
    if want["kind"] == "image":
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            # python-pptx has no image part for WebP, so the content type is read off the part the picture points at.
            content_type = shape.part.related_part(shape._element.blip_rId).content_type
            if content_type != want["image"]["mime"]:
                problems.append(f"{tag}: picture {content_type} for {want['image']['mime']}")
            crop = want["image"].get("crop")
            if crop:
                for side, v in (("left", crop["l"]), ("top", crop["t"]), ("right", crop["r"]), ("bottom", crop["b"])):
                    if abs(getattr(shape, f"crop_{side}") - v) > 0.001:
                        problems.append(f"{tag}: crop {side} {getattr(shape, f'crop_{side}')} for {v}")
    if want["kind"] == "group":
        kids = list(shape.shapes)
        if len(kids) != len(want["children"]):
            problems.append(f"{tag}: {len(kids)} children for {len(want['children'])}")
        else:
            for i, (k, w) in enumerate(zip(kids, want["children"])):
                compare(k, w, f"{tag}.{i + 1}", emu_per_px, problems)


def workbook_problems(blob, chart, tag, problems):
    """The workbook behind a chart holds what the chart shows, if openpyxl is here to read it."""
    try:
        import io

        import openpyxl
    except ImportError:
        return
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(io.BytesIO(blob), data_only=True).active
    if [c.value for c in ws[1]][1 : 1 + len(chart["series"])] != [s["name"] for s in chart["series"]]:
        problems.append(f"{tag}: the workbook's series names {[c.value for c in ws[1]]}")
    for k, s in enumerate(chart["series"]):
        got = [float(ws.cell(row=r + 2, column=k + 2).value) for r in range(len(s["values"]))]
        if got != [float(v) for v in s["values"]]:
            problems.append(f"{tag}: the workbook's {s['name']} is {got}")
    if chart["chart"] != "scatter":
        cats = [str(ws.cell(row=r + 2, column=1).value) for r in range(len(chart["categories"]))]
        if cats != chart["categories"]:
            problems.append(f"{tag}: the workbook's categories {cats}")


# Elements that are a choice and not a sequence: their children may come in any order (a z-order).
CHOICE = {"p:spTree", "p:grpSp"}
_ORDER = None


def _name(e):
    ns, _, local = e.tag[1:].partition("}")
    prefix = {
        "http://schemas.openxmlformats.org/drawingml/2006/main": "a",
        "http://schemas.openxmlformats.org/presentationml/2006/main": "p",
        "http://schemas.openxmlformats.org/drawingml/2006/chart": "c",
        "http://schemas.openxmlformats.org/markup-compatibility/2006": "mc",
    }.get(ns, ns)
    return f"{prefix}:{local}"


def _sequences(z):
    for n in z.namelist():
        if n.endswith(".xml") and n.startswith("ppt/"):
            try:
                root = ET.fromstring(z.read(n))
            except ET.ParseError:
                continue
            for e in root.iter():
                kids = [_name(c) for c in e if isinstance(c.tag, str)]
                if len(kids) > 1:
                    yield n, _name(e), kids


def reference_order(dir_: Path):
    """Which sibling came before which in files PowerPoint's own template and python-pptx wrote."""
    import pptx

    global _ORDER
    if _ORDER is not None:
        return _ORDER
    before = {}
    refs = [Path(pptx.__file__).parent / "templates" / "default.pptx", *sorted(dir_.glob("../*python-pptx-*.pptx")), *sorted(dir_.parent.glob("*python-pptx-*.pptx"))]
    for ref in refs:
        if not ref.exists():
            continue
        with zipfile.ZipFile(ref) as z:
            for _, parent, kids in _sequences(z):
                for i, a in enumerate(kids):
                    for b in kids[i + 1 :]:
                        if a != b:
                            before.setdefault((parent, a), set()).add(b)
    _ORDER = before
    return before


def order_problems(path: Path):
    """No part puts two siblings the other way round from every file PowerPoint's template and python-pptx wrote."""
    before = reference_order(path.parent)
    problems = set()
    with zipfile.ZipFile(path) as z:
        for n, parent, kids in _sequences(z):
            if parent in CHOICE:
                continue
            for i, a in enumerate(kids):
                for b in kids[i + 1 :]:
                    if a != b and a in before.get((parent, b), set()) and b not in before.get((parent, a), set()):
                        problems.add(f"{n}: {parent} has {a} before {b}")
    return sorted(problems)


def check(path: Path):
    want = json.loads(path.with_suffix(".check.json").read_text(encoding="utf-8"))
    problems = package_problems(path) + order_problems(path)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        prs = Presentation(str(path))
    if (prs.slide_width, prs.slide_height) != (want["size"]["cx"], want["size"]["cy"]):
        problems.append(f"size {prs.slide_width}x{prs.slide_height}")
    layouts = [l.name for l in prs.slide_layouts]
    if layouts != want["layouts"]:
        problems.append(f"layouts {layouts} != {want['layouts']}")
    slides = list(prs.slides)
    if len(slides) != len(want["slides"]):
        problems.append(f"{len(slides)} slides != {len(want['slides'])}")
        return problems
    for i, (slide, w) in enumerate(zip(slides, want["slides"])):
        tag = f"slide {i + 1}"
        if slide.slide_layout.name != w["layout"]:
            problems.append(f"{tag}: layout {slide.slide_layout.name!r} != {w['layout']!r}")
        if (slide._element.get("show") == "0") != w["hidden"]:
            problems.append(f"{tag}: hidden")
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else None
        if (notes or None) != w["notes"]:
            problems.append(f"{tag}: notes {notes!r} != {w['notes']!r}")
        shapes = list(slide.shapes)
        if len(shapes) != len(w["elements"]):
            problems.append(f"{tag}: {len(shapes)} shapes != {len(w['elements'])}")
            continue
        for k, (shape, e) in enumerate(zip(shapes, w["elements"])):
            compare(shape, e, f"{tag}, element {k + 1}", want["emuPerPx"], problems)
    return problems


def main():
    out = Path(sys.argv[1])
    report = []
    for p in sorted(out.glob("*.pptx")):
        try:
            problems = check(p)
        except Exception as e:  # noqa: BLE001
            problems = [f"{type(e).__name__}: {e}"]
        report.append({"file": p.name, "problems": problems})
    print(json.dumps(report))
    sys.exit(1 if any(r["problems"] for r in report) else 0)


if __name__ == "__main__":
    main()
