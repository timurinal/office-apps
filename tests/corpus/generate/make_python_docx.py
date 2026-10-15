"""
The python-docx half of the corpus: ten documents, each covering one family of features.

    python3 -m venv tests/corpus/.venv
    tests/corpus/.venv/bin/pip install python-docx
    tests/corpus/.venv/bin/python tests/corpus/generate/make_python_docx.py

python-docx writes what Word's own templates write (styles.xml, numbering.xml, theme), so
these are closer to a person's documents than anything the `docx` package makes. The
files are committed; this script is how they were made and how to change one.
"""

import base64
import copy
import struct
import sys
import zlib
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "docx"


def png(width: int, height: int, rgb) -> bytes:
    """A gradient PNG, made by hand so no imaging library is needed."""
    rows = bytearray()
    for y in range(height):
        rows.append(0)
        for x in range(width):
            rows.extend((min(255, rgb[0] + x), min(255, rgb[1] + y), rgb[2]))

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(bytes(rows)))
        + chunk(b"IEND", b"")
    )


JPEG_1PX = base64.b64decode(
    "/9j/4AAQSkZJRgABAQEASABIAAD/2wBDAP//////////////////////////////////////////////////////////////////////////////////////wgALCAABAAEBAREA/8QAFBABAAAAAAAAAAAAAAAAAAAAAP/aAAgBAQABPxA="
)
GIF_1PX = base64.b64decode("R0lGODlhAQABAIAAAAUEBAAAACwAAAAAAQABAAACAkQBADs=")


def save(doc: Document, name: str) -> None:
    doc.core_properties.author = "Corpus"
    doc.core_properties.created = datetime(2026, 1, 1, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    doc.save(OUT / name)
    print("wrote", name)


def shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_hyperlink(paragraph, url: str, text: str):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rPr.append(style)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_field(paragraph, instruction: str, placeholder: str = "1"):
    for kind in ("begin", None, "separate", "text", "end"):
        run = paragraph.add_run()
        if kind in ("begin", "separate", "end"):
            c = OxmlElement("w:fldChar")
            c.set(qn("w:fldCharType"), kind)
            run._r.append(c)
        elif kind is None:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = f" {instruction} "
            run._r.append(t)
        else:
            run.text = placeholder


# ── 01 ───────────────────────────────────────────────────────────────────────
def formatting():
    d = Document()
    d.add_heading("Character and paragraph formatting", level=1)
    p = d.add_paragraph("Plain, ")
    for text, attr in [("bold", "bold"), ("italic", "italic"), ("underlined", "underline"), ("struck through", "strike")]:
        r = p.add_run(text)
        setattr(r.font, attr, True) if attr != "strike" else setattr(r.font, "strike", True)
        p.add_run(", ")
    r = p.add_run("subscript")
    r.font.subscript = True
    p.add_run(" and ")
    r = p.add_run("superscript")
    r.font.superscript = True
    p.add_run(".")

    p = d.add_paragraph()
    r = p.add_run("Georgia at 16 pt in dark red")
    r.font.name = "Georgia"
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x8E, 0x1F, 0x1F)
    p.add_run(" then ")
    r = p.add_run("highlighted yellow")
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.add_run(" then ")
    r = p.add_run("Courier New 9 pt")
    r.font.name = "Courier New"
    r.font.size = Pt(9)

    for align, label in [(WD_ALIGN_PARAGRAPH.LEFT, "Left"), (WD_ALIGN_PARAGRAPH.CENTER, "Centred"), (WD_ALIGN_PARAGRAPH.RIGHT, "Right"), (WD_ALIGN_PARAGRAPH.JUSTIFY, "Justified")]:
        p = d.add_paragraph(f"{label} paragraph. " + "The quick brown fox jumps over the lazy dog. " * (4 if label == "Justified" else 1))
        p.alignment = align

    p = d.add_paragraph("Spaced paragraph: 18 pt before, 6 pt after, double line spacing.")
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 2.0
    p = d.add_paragraph("Indented by one inch.")
    p.paragraph_format.left_indent = Inches(1)
    p = d.add_paragraph("Indented by two inches, 1.5 line spacing.")
    p.paragraph_format.left_indent = Inches(2)
    p.paragraph_format.line_spacing = 1.5
    save(d, "01-formatting.docx")


# ── 02 ───────────────────────────────────────────────────────────────────────
def lists():
    d = Document()
    d.add_heading("Lists", level=1)
    d.add_paragraph("A bulleted list:")
    d.add_paragraph("First bullet", style="List Bullet")
    d.add_paragraph("Second bullet", style="List Bullet")
    d.add_paragraph("A nested bullet", style="List Bullet 2")
    d.add_paragraph("Another nested bullet", style="List Bullet 2")
    d.add_paragraph("Third bullet", style="List Bullet")
    d.add_paragraph("A numbered list:")
    d.add_paragraph("Step one", style="List Number")
    d.add_paragraph("Step two", style="List Number")
    d.add_paragraph("A nested step", style="List Number 2")
    d.add_paragraph("Step three", style="List Number")
    d.add_paragraph("Text between the lists.")
    d.add_paragraph("Bullets again", style="List Bullet")
    d.add_paragraph("And again", style="List Bullet")
    save(d, "02-lists.docx")


# ── 03 ───────────────────────────────────────────────────────────────────────
def tables():
    d = Document()
    d.add_heading("Tables", level=1)
    d.add_paragraph("A grid with a header row and shading:")
    t = d.add_table(rows=3, cols=3, style="Table Grid")
    for i, h in enumerate(["Name", "Role", "Start"]):
        t.cell(0, i).text = h
        shade(t.cell(0, i), "DBE7FD")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    for r, row in enumerate([("Ada Rowe", "Editor", "2024"), ("Ben Okoro", "Reviewer", "2025")], start=1):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.add_paragraph()
    d.add_paragraph("Merged cells, across and down:")
    m = d.add_table(rows=4, cols=4, style="Table Grid")
    m.cell(0, 0).merge(m.cell(0, 3)).text = "Across all four columns"
    m.cell(1, 0).merge(m.cell(3, 0)).text = "Down three rows"
    m.cell(1, 1).merge(m.cell(1, 2)).text = "Two across"
    m.cell(1, 3).text = "x"
    m.cell(2, 1).text = "a"
    m.cell(2, 2).text = "b"
    m.cell(2, 3).text = "c"
    m.cell(3, 1).text = "d"
    m.cell(3, 2).text = "e"
    m.cell(3, 3).text = "f"
    d.add_paragraph()
    d.add_paragraph("A table inside a table cell:")
    outer = d.add_table(rows=1, cols=2, style="Table Grid")
    outer.cell(0, 0).text = "Outside"
    inner = outer.cell(0, 1).add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            inner.cell(r, c).text = f"in {r},{c}"
    outer.cell(0, 1).add_paragraph("After the inner table.")
    d.add_paragraph("The end.")
    save(d, "03-tables.docx")


# ── 04 ───────────────────────────────────────────────────────────────────────
def images():
    d = Document()
    d.add_heading("Pictures", level=1)
    p = d.add_paragraph("A picture in the middle of a sentence: before ")
    p.add_run().add_picture(BytesIO(png(120, 60, (20, 80, 200))), width=Inches(1.2))
    p.add_run(" and after.")
    cap = d.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run().add_picture(BytesIO(png(200, 100, (200, 60, 20))), width=Inches(3))
    shape = cap.runs[0]._r.xpath(".//wp:docPr")[0]
    shape.set("descr", "A centred orange gradient")
    d.add_paragraph("A JPEG and a GIF:")
    q = d.add_paragraph()
    q.add_run().add_picture(BytesIO(JPEG_1PX), width=Inches(0.5))
    q.add_run().add_picture(BytesIO(GIF_1PX), width=Inches(0.5))
    d.add_paragraph("The same picture twice:")
    same = png(80, 80, (10, 160, 90))
    d.add_paragraph().add_run().add_picture(BytesIO(same), width=Inches(1))
    d.add_paragraph().add_run().add_picture(BytesIO(same), width=Inches(2))
    save(d, "04-images.docx")


# ── 05 ───────────────────────────────────────────────────────────────────────
def comments():
    d = Document()
    d.add_heading("Comments", level=1)
    p = d.add_paragraph("The first sentence has a comment on this phrase and then carries on. ")
    p.add_run("A second sentence, partly bold, ").bold = True
    last = p.add_run("with a comment spanning two runs.")
    first_run = p.runs[0]
    d.add_comment(runs=first_run, text="Is this the right word?", author="Ben Okoro", initials="BO")
    d.add_comment(runs=[p.runs[1], last], text="Check the figure.\nSee the appendix.", author="Ada Rowe", initials="AR")
    d.add_paragraph("A paragraph with no comment.")
    save(d, "05-comments.docx")


# ── 06 ───────────────────────────────────────────────────────────────────────
def page_setup():
    d = Document()
    s = d.sections[0]
    s.page_width = Inches(11)
    s.page_height = Inches(8.5)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.left_margin = Inches(0.75)
    s.right_margin = Inches(0.75)
    s.top_margin = Inches(1)
    s.bottom_margin = Inches(0.9)
    s.header_distance = Inches(0.4)
    s.footer_distance = Inches(0.4)
    s.different_first_page_header_footer = True
    s.first_page_header.paragraphs[0].text = "Title page header"
    s.header.paragraphs[0].text = "Quarterly report"
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ")
    add_field(fp, "PAGE")
    fp.add_run(" of ")
    add_field(fp, "NUMPAGES", "9")
    s.first_page_footer.paragraphs[0].text = "First page footer"
    d.add_heading("Landscape, Letter, custom margins", level=1)
    for i in range(30):
        d.add_paragraph(f"Paragraph {i + 1}. " + "Words that fill a line and then some more. " * 4)
    save(d, "06-page-setup.docx")


# ── 07 ───────────────────────────────────────────────────────────────────────
def styles():
    d = Document()
    d.add_paragraph("A title", style="Title")
    d.add_paragraph("A subtitle", style="Subtitle")
    for level in range(1, 5):
        d.add_heading(f"Heading {level}", level=level)
    d.add_paragraph("Normal paragraph after the headings.")
    d.add_paragraph("A quotation that is set apart.", style="Quote")
    d.add_paragraph("An intense quotation.", style="Intense Quote")
    d.add_paragraph("No spacing, line one.", style="No Spacing")
    d.add_paragraph("No spacing, line two.", style="No Spacing")

    from docx.enum.style import WD_STYLE_TYPE

    body2 = d.styles.add_style("Body Text Two", WD_STYLE_TYPE.PARAGRAPH)
    body2.base_style = d.styles["Normal"]
    body2.font.bold = True
    body2.font.color.rgb = RGBColor(0x1F, 0x4E, 0xA8)
    body2.font.size = Pt(13)
    body2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("A custom style: bold, blue, centred, 13 pt.", style="Body Text Two")
    child = d.styles.add_style("Body Text Three", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = body2
    child.font.italic = True
    d.add_paragraph("Based on the one above, and italic as well.", style="Body Text Three")
    chap = d.styles.add_style("Chapter Heading", WD_STYLE_TYPE.PARAGRAPH)
    chap.base_style = d.styles["Heading 2"]
    d.add_paragraph("Based on Heading 2", style="Chapter Heading")
    emph = d.styles.add_style("Brand", WD_STYLE_TYPE.CHARACTER)
    emph.font.bold = True
    emph.font.color.rgb = RGBColor(0xC0, 0x50, 0x1C)
    p = d.add_paragraph("A character style: ")
    p.add_run("Acme", style="Brand")
    save(d, "07-styles.docx")


# ── 08 ───────────────────────────────────────────────────────────────────────
def breaks():
    d = Document()
    d.add_paragraph("Page one.")
    d.add_page_break()
    d.add_paragraph("Page two, after a page break.")
    p = d.add_paragraph("Page three, because this paragraph asks for a break before it.")
    p.paragraph_format.page_break_before = True
    d.add_section(WD_SECTION.NEW_PAGE)
    d.add_paragraph("Page four, a new section.")
    d.add_section(WD_SECTION.CONTINUOUS)
    d.add_paragraph("A continuous section on the same page.")
    s = d.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    d.add_paragraph("A landscape section.")
    save(d, "08-breaks-sections.docx")


# ── 09 ───────────────────────────────────────────────────────────────────────
def unicode_text():
    d = Document()
    d.add_heading("Many scripts", level=1)
    for label, text in [
        ("Latin", "Árvíztűrő tükörfúrógép — naïve façade, Ærøskøbing, Łódź."),
        ("Cyrillic", "Съешь же ещё этих мягких французских булок, да выпей чаю."),
        ("Greek", "Ξεσκεπάζω την ψυχοφθόρα βδελυγμία."),
        ("CJK", "日本語のテキスト、中文文本、한국어 텍스트。"),
        ("Arabic", "نص عربي بسيط يُقرأ من اليمين إلى اليسار."),
        ("Hebrew", "טקסט עברי פשוט הנקרא מימין לשמאל."),
        ("Emoji", "A few emoji: 🙂 🚀 ✅ and a flag 🇳🇴."),
        ("Symbols", "© ® ™ § ¶ † ‡ • … – — ‘quoted’ “quoted” € £ ¥ ½ ² ³ ≠ ≤ ≥ ∞"),
    ]:
        p = d.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(text)
    p = d.add_paragraph("Tabs and breaks:\tafter a tab\tand another")
    run = p.add_run()
    run.add_break(WD_BREAK.LINE)
    run.add_text("after a line break, with a non-breaking space and a non‑breaking hyphen.")
    d.add_paragraph("")
    d.add_paragraph("Whitespace   kept   inside   the   line,   and   a   very   long   word:   " + "x" * 120)
    save(d, "09-unicode.docx")


# ── 10 ───────────────────────────────────────────────────────────────────────
def large():
    d = Document()
    d.add_paragraph("A long document", style="Title")
    for chapter in range(1, 13):
        d.add_heading(f"Chapter {chapter}", level=1)
        for section in range(1, 4):
            d.add_heading(f"Section {chapter}.{section}", level=2)
            for i in range(6):
                d.add_paragraph(
                    f"Paragraph {i + 1} of section {chapter}.{section}. "
                    + "The editor has to lay this out over many pages, keep every character, and stay quick while it does. " * 3
                )
            if section == 2:
                t = d.add_table(rows=5, cols=4, style="Table Grid")
                for r in range(5):
                    for c in range(4):
                        t.cell(r, c).text = f"R{r + 1}C{c + 1}"
                d.add_paragraph()
    save(d, "10-large.docx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    only = set(sys.argv[1:])
    for fn in (formatting, lists, tables, images, comments, page_setup, styles, breaks, unicode_text, large):
        if not only or fn.__name__ in only:
            fn()
